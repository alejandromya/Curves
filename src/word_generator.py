# src/word_generator.py
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import io
import numpy as np

def _estilizar_tabla_docx(table):
    """
    Estiliza tabla: header gris y centrar texto.
    """
    # aplicar estilo base
    try:
        table.style = "Table Grid"
    except Exception:
        pass

    # Header gris
    hdr = table.rows[0]
    for cell in hdr.cells:
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), "D9D9D9")  # gris suave
        cell._tc.get_or_add_tcPr().append(shading)

    # Centrar todo el texto
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = 1  # centered

def generar_tabla_ciclos_word(detalles_ciclos, df, fuerza_max, deformacion_max):
    """
    Genera headers y fila (valores) para la tabla, con formato idéntico al PDF.
    Todos los números a 2 decimales.
    """

    if not detalles_ciclos:
        headers = ["0", "F2mm", "F3mm", f"Fmax ({fuerza_max:.2f} N)", "Dciclico"]
        fila = ["—", "—", "—", f"{deformacion_max:.2f}", "—"]
        return headers, fila

    # claves ordenadas de ciclos (p. ej. 1,2,... o 100,200,...)
    claves = sorted(detalles_ciclos.keys())
    primer_key = claves[0]
    ultimo_key = claves[-1]
    primer_ciclo = detalles_ciclos[primer_key]
    ultimo_ciclo = detalles_ciclos[ultimo_key]

    # determinar última centena disponible
    ult_centena = (ultimo_key // 100) * 100

    # columnas de ciclos: 0,100,200,...,ult_centena
    ciclos_col = [0] + [c for c in range(100, ult_centena + 1, 100)]

    # asegurar que la última centena esté presente (ya lo está por construcción)
    # construir headers (última columna Fmax muestra fuerza en header)
    headers = [str(c) for c in ciclos_col] + ["F2mm", "F3mm", f"Fmax ({fuerza_max:.2f} N)", "Dciclico"]

    # construir fila de valores (solo HIGHs: 0 -> deform_high_start, demás -> deform_high_end)
    fila = []
    for c in ciclos_col:
        if c == 0:
            # ciclo 0 := HIGH inicio del primer ciclo real (primer_key)
            # usar deform_high_start del primer ciclo
            val = primer_ciclo.get("deform_high_start", None)
            if val is None:
                fila.append("—")
            else:
                fila.append(f"{val:.2f}")
        else:
            # si existe el ciclo con esa clave, tomar deform_high_end, si no → "—"
            if c in detalles_ciclos:
                val = detalles_ciclos[c].get("deform_high_end", None)
                fila.append(f"{val:.2f}" if val is not None else "—")
            else:
                fila.append("—")

    # calcular low del último ciclo (base para F2mm/F3mm y para Fmax relativo)
    deform_low_last = ultimo_ciclo.get("deform_low", None)
    if deform_low_last is None:
        # fallback: intentar usar valle_idx para extraer deformación desde df
        if "valle_idx" in ultimo_ciclo:
            idx = ultimo_ciclo["valle_idx"]
            deform_low_last = float(df.loc[idx, "Deformacion"])
        else:
            deform_low_last = None

    # F2mm y F3mm (buscamos el punto más cercano a deform_low_last + 2/3 mm)
    if deform_low_last is not None:
        arr_def = df["Deformacion"].values
        # índice más cercano
        i2 = int(np.abs(arr_def - (deform_low_last + 2.0)).argmin())
        i3 = int(np.abs(arr_def - (deform_low_last + 3.0)).argmin())

        f2_val = float(df.loc[i2, "Fuerza"])
        f3_val = float(df.loc[i3, "Fuerza"])
        fila.append(f"{f2_val:.2f}")
        fila.append(f"{f3_val:.2f}")
    else:
        fila.append("—")
        fila.append("—")

    # Fmax: según PDF actual mostramos el desplazamiento relativo al low del último ciclo
    if deform_low_last is not None:
        df_max_rel = deformacion_max - deform_low_last
        fila.append(f"{df_max_rel:.2f}")
    else:
        fila.append(f"{deformacion_max:.2f}")  # fallback: mostrar deformacion_max absoluta

    # Dciclico = High final último ciclo - High inicio primer ciclo
    d_high_last = ultimo_ciclo.get("deform_high_end", None)
    d_high_start0 = primer_ciclo.get("deform_high_start", None)
    if (d_high_last is not None) and (d_high_start0 is not None):
        dc = d_high_last - d_high_start0
        fila.append(f"{dc:.2f}")
    else:
        fila.append("—")

    return headers, fila

def generar_word_unico(bloques, output_docx="INFORME_FINAL.docx"):
    """
    Genera un documento Word con los mismos contenidos/tablas que el PDF.
    Cada bloque debe contener:
      - 'titulo' (str)
      - 'ciclos' (dict)
      - 'df' (pandas.DataFrame)
      - 'grafico' (BytesIO)
      - 'fuerza_max' (float)
      - 'deformacion_max' (float)
    """

    doc = Document()

    for bloque in bloques:
        # Título
        doc.add_heading(f"{bloque['titulo']}", level=1)

        # Datos generales (force y desplazamiento absolutos, 2 decimales)
        doc.add_paragraph(f"Fuerza máxima: {bloque['fuerza_max']:.2f} N")
        doc.add_paragraph(f"Desplazamiento en fuerza máxima: {bloque['deformacion_max']:.2f} mm")
        doc.add_paragraph("")  # espacio

        # Generar tabla (headers y fila)
        headers, fila = generar_tabla_ciclos_word(bloque["ciclos"], bloque["df"],
                                                 bloque["fuerza_max"], bloque["deformacion_max"])

        # Insertar tabla en Word
        table = doc.add_table(rows=1, cols=len(headers))
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = str(h)

        row_cells = table.add_row().cells
        for i, value in enumerate(fila):
            row_cells[i].text = str(value)

        # Estilizar
        _estilizar_tabla_docx(table)

        doc.add_paragraph("")  # espacio

        # Imagen: bloque["grafico"] es BytesIO -> añadir
        graf = bloque.get("grafico", None)
        if graf is not None:
            # asegurarse de estar al inicio
            try:
                graf.seek(0)
                doc.add_picture(graf, width=Inches(5))
            except Exception:
                # si no funciona directo (p.ej. str), intentar BytesIO desde getvalue()
                try:
                    img_bytes = io.BytesIO(graf.getvalue())
                    img_bytes.seek(0)
                    doc.add_picture(img_bytes, width=Inches(5))
                except Exception:
                    pass

        doc.add_page_break()

    doc.save(output_docx)
    return output_docx
